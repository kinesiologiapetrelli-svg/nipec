# Genera CONTATTI-CAMPAGNA.xlsx: python3 CAMPAGNA-MATERIALI/_foglio.py
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

QUI = os.path.dirname(os.path.abspath(__file__))
NAVY, BLU, GIALLO, TENUE = '123A63', 'E6EFF7', 'FFF6D6', '5A6B79'
CODICI = ['R1', 'R2', 'R3', 'R4', 'R5', 'R6', 'R7', 'R8', 'IMG-N', 'IMG-A', 'IMG-P', 'search', 'altro']
N = 500  # righe pronte

wb = Workbook()
c = wb.active
c.title = 'CONTATTI'
testata = ['Data', 'Canale', 'Codice annuncio', 'Nome', 'Centro', 'Città', 'Telefono', 'Laser oggi',
           'Agenda', 'Qualificata?', 'Chiamata fissata?', 'Data chiamata', 'Esito', 'Note']
larg = [12, 15, 15, 20, 22, 16, 16, 16, 14, 13, 16, 14, 20, 40]
c['A1'] = 'CONTATTI · campagna noleggio laser NIPEC'
c['A1'].font = Font(bold=True, size=14, color=NAVY)
c['A2'] = 'Una riga per ogni chat o modulo. Le caselle gialle hanno la tendina. Il codice è quello tra parentesi nel primo messaggio: [rif: R3].'
c['A2'].font = Font(italic=True, color=TENUE)
for i, (h, w) in enumerate(zip(testata, larg), 1):
    x = c.cell(row=4, column=i, value=h)
    x.font = Font(bold=True, color='FFFFFF')
    x.fill = PatternFill('solid', fgColor=NAVY)
    x.alignment = Alignment(vertical='center', wrap_text=True)
    c.column_dimensions[x.column_letter].width = w
c.row_dimensions[4].height = 30
c.freeze_panes = 'E5'

def tendina(col, voci):
    dv = DataValidation(type='list', formula1='"%s"' % ','.join(voci), allow_blank=True)
    dv.error, dv.errorTitle = 'Scegli una voce della tendina', 'Valore non previsto'
    c.add_data_validation(dv)
    dv.add('%s5:%s%d' % (col, col, N + 4))
    for r in range(5, N + 5):
        c['%s%d' % (col, r)].fill = PatternFill('solid', fgColor=GIALLO)

tendina('B', ['WhatsApp', 'Modulo sito', 'Modulo Meta', 'Google', 'Telefono', 'Instagram DM'])
tendina('C', CODICI)
tendina('H', ['Nessuno', 'Vecchio', 'Recente'])
tendina('I', ['Quaderno', 'Telefono', 'Gestionale'])
tendina('J', ['Sì', 'No'])
tendina('K', ['Sì', 'No'])
tendina('M', ['In corso', 'Noleggio firmato', 'Acquisto firmato', 'Non interessata', 'Non risponde'])
for r in range(5, N + 5):
    c['A%d' % r].number_format = 'DD/MM/YYYY'
    c['L%d' % r].number_format = 'DD/MM/YYYY'
    c['G%d' % r].number_format = '@'
c.conditional_formatting.add('A5:N%d' % (N + 4), FormulaRule(formula=['OR($M5="Noleggio firmato",$M5="Acquisto firmato")'], fill=PatternFill('solid', fgColor='E2F1EA')))
c.conditional_formatting.add('A5:N%d' % (N + 4), FormulaRule(formula=['AND($A5<>"",$K5<>"Sì",$M5<>"Non interessata",TODAY()-$A5>2)'], fill=PatternFill('solid', fgColor='FAEAE6')))

s = wb.create_sheet('NUMERI')
s['A1'] = 'NUMERI · per annuncio'
s['A1'].font = Font(bold=True, size=14, color=NAVY)
s['A2'] = 'Si calcolano da soli dal foglio CONTATTI. Scrivi a mano solo la colonna gialla «Speso €», copiandola da Gestione inserzioni.'
s['A2'].font = Font(italic=True, color=TENUE)
h2 = ['Codice', 'Speso €', 'Contatti', 'Qualificate', 'Chiamate fissate', 'Contratti', '€ per contatto', '€ per chiamata', '€ per contratto']
for i, h in enumerate(h2, 1):
    x = s.cell(row=4, column=i, value=h)
    x.font = Font(bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor=NAVY)
    x.alignment = Alignment(wrap_text=True, vertical='center')
    s.column_dimensions[x.column_letter].width = 15
s.row_dimensions[4].height = 30
R = 'CONTATTI!$C$5:$C$%d' % (N + 4)
for j, cod in enumerate(CODICI):
    r = 5 + j
    s['A%d' % r] = cod
    s['B%d' % r].fill = PatternFill('solid', fgColor=GIALLO)
    s['B%d' % r].number_format = '#,##0.00 €'
    s['C%d' % r] = '=COUNTIF(%s,A%d)' % (R, r)
    s['D%d' % r] = '=COUNTIFS(%s,A%d,CONTATTI!$J$5:$J$%d,"Sì")' % (R, r, N + 4)
    s['E%d' % r] = '=COUNTIFS(%s,A%d,CONTATTI!$K$5:$K$%d,"Sì")' % (R, r, N + 4)
    s['F%d' % r] = '=COUNTIFS(%s,A%d,CONTATTI!$M$5:$M$%d,"Noleggio firmato")+COUNTIFS(%s,A%d,CONTATTI!$M$5:$M$%d,"Acquisto firmato")' % (R, r, N + 4, R, r, N + 4)
    for col, den in (('G', 'C'), ('H', 'E'), ('I', 'F')):
        s['%s%d' % (col, r)] = '=IF(OR(B{r}="",{d}{r}=0),"",B{r}/{d}{r})'.format(r=r, d=den)
        s['%s%d' % (col, r)].number_format = '#,##0 €'
t = 5 + len(CODICI)
s['A%d' % t] = 'TOTALE'
s['A%d' % t].font = Font(bold=True)
for col in 'BCDEF':
    s['%s%d' % (col, t)] = '=SUM(%s5:%s%d)' % (col, col, t - 1)
    s['%s%d' % (col, t)].font = Font(bold=True)
s['B%d' % t].number_format = '#,##0.00 €'
for col, den in (('G', 'C'), ('H', 'E'), ('I', 'F')):
    s['%s%d' % (col, t)] = '=IF({d}{r}=0,"",B{r}/{d}{r})'.format(r=t, d=den)
    s['%s%d' % (col, t)].number_format = '#,##0 €'
    s['%s%d' % (col, t)].font = Font(bold=True)
s['A%d' % (t + 2)] = 'Le soglie di partenza: contatto sotto i 15 €, chiamata fissata sotto i 150 €. Dopo due settimane valgono i vostri numeri.'
s['A%d' % (t + 2)].font = Font(italic=True, color=TENUE)
wb.save(os.path.join(QUI, 'CONTATTI-CAMPAGNA.xlsx'))
print('ok')
