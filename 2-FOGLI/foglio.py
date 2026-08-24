# -*- coding: utf-8 -*-
"""Costruisce il Cruscotto Clienti · laser. Le formule vengono da formule.py:
qui non se ne scrive nessuna a mano."""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from formule import F, perExcel, CELLE, rendi

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

NAVY, BLU, CHIARO, INK = "FF123A63", "FF1C6EA4", "FF5AA6D6", "FF17222E"
GIALLO  = PatternFill("solid", fgColor="FFFDF3D0")   # lo scrive lei
CALC    = PatternFill("solid", fgColor="FFF2F5F8")   # si calcola da solo
TESTATA = PatternFill("solid", fgColor=NAVY)
BANDA   = PatternFill("solid", fgColor="FFE6EFF7")
ROSSO_F = PatternFill("solid", fgColor="FFFAEAE6"); ROSSO_T = Font(name="Arial", size=10, color="FFA3432E", bold=True)
GIALL_F = PatternFill("solid", fgColor="FFF8EFDB"); GIALL_T = Font(name="Arial", size=10, color="FF8A6208", bold=True)
VERDE_F = PatternFill("solid", fgColor="FFE2F1EA"); VERDE_T = Font(name="Arial", size=10, color="FF1B6E52", bold=True)
SOTTILE = Border(bottom=Side(style="thin", color="FFD9E2E9"))

A   = lambda **k: Font(name="Arial", size=10, **k)
TIT = lambda s=16: Font(name="Arial", size=s, bold=True, color=NAVY)

STATI  = ["Da contattare", "In ciclo", "Mantenimento", "Sospesa", "Persa"]
AVVISI = ["CONFERMA DOMANI", "CHIAMA PER CONFERMA", "APPUNTAMENTO SALTATO", "PRENDI APPUNTAMENTO",
          "FUORI FINESTRA", "ULTIMA SEDUTA", "RICHIAMO MANTENIMENTO", "RECUPERA CON INCENTIVO"]

# colonna, intestazione, larghezza, chi la scrive, formato
COLONNE = [
 ("A", "Nome",                  24, "lei",     None),
 ("B", "COSA FARE OGGI",        26, "formula", None),
 ("C", "Prossimo appuntamento", 15, "lei",     "dd/mm/yyyy"),
 ("D", "Sedute fatte",          10, "lei",     "0"),
 ("E", "Stato",                 15, "lei",     None),
 ("F", "Promemoria",            34, "formula", None),
 ("G", "Telefono",              14, "lei",     "@"),
 ("H", "Zona",                  14, "lei",     None),
 ("I", "Ultima seduta",         13, "lei",     "dd/mm/yyyy"),
 ("J", "Giorni dall'ultima",    11, "formula", "0"),
 ("K", "Giorni al prossimo",    11, "formula", "0"),
 ("L", "Sedute",                 8, "formula", None),
 ("M", "Offerta da usare",      20, "formula", None),
 ("N", "Prezzo a seduta €",     12, "lei",     '#,##0 "€"'),
 ("O", "Valore lasciato €",     13, "formula", '#,##0 "€"'),
 ("P", "Ancora da fare €",      13, "formula", '#,##0 "€"'),
 ("Q", "Inserita il",           12, "lei",     "dd/mm/yyyy"),
 ("R", "Note",                  30, "lei",     None),
 ("S", "avviso (servizio)",     22, "formula", None),
 ("T", "ordine (servizio)",     14, "formula", "0"),
]
PRIMA, ULTIMA = 5, 304


def perRiga(chiave, riga):
    """La formula di formule.py portata dalla riga 5 alla riga giusta."""
    f = perExcel(chiave)
    if riga == 5:
        return f
    for c in "ABCDEFGHIJKLMNOPQRST":
        f = f.replace("$%s5" % c, "$%s%d" % (c, riga))
    return f


def titolo(ws, testo, sotto, largo="H"):
    ws["B2"] = testo; ws["B2"].font = TIT()
    ws["B3"] = sotto; ws["B3"].font = A(color="FF5A6B79")
    ws.merge_cells("B3:%s3" % largo)


def leggimi(wb):
    ws = wb.create_sheet("LEGGIMI")
    ws.sheet_view.showGridLines = False
    titolo(ws, "Cruscotto Clienti · laser", "NIPEC · Tecnologie Estetiche", "H")
    righe = [
      ("", ""),
      ("A COSA SERVE", ""),
      ("", "Dice ogni mattina chi richiamare e in che ordine. Si guarda una colonna sola:"),
      ("", "COSA FARE OGGI. Si ordina per quella e si lavora dall'alto."),
      ("", ""),
      ("LE TRE REGOLE", ""),
      ("Giallo", "lo scrivi tu. Sono le uniche celle da toccare."),
      ("Grigio", "si calcola da solo. Se ci scrivi dentro, quella riga smette di avvisare."),
      ("Riga 5", "è un esempio. Seleziona le celle gialle e cancellale quando cominci."),
      ("", ""),
      ("IL GESTO DI OGNI GIORNO", ""),
      ("", "La cliente viene: alzi Sedute fatte di uno, metti Ultima seduta a oggi,"),
      ("", "e fissi il Prossimo appuntamento PRIMA che esca dalla porta."),
      ("", "Se fai le prime due e salti la terza, domattina te la ritrovi in cima."),
      ("", ""),
      ("I FOGLI", ""),
      ("CLIENTI", "la rubrica. Una riga per persona, o una per zona se i cicli sono separati."),
      ("OGGI", "solo chi va chiamato, già in ordine. È quello da aprire la mattina."),
      ("RIEPILOGO", "i numeri del centro. Da qui il master NIPEC legge sei celle: non spostarle."),
      ("INCENTIVI", "le tre offerte. Le rinomini qui e il cruscotto le segue."),
      ("IMPOSTAZIONI", "i sei numeri dell'orologio del laser. Cambi lì, cambia tutto."),
      ("", ""),
      ("LA REGOLA D'ORO", ""),
      ("", "Nessuna esce senza il prossimo appuntamento."),
      ("", "Sotto il 90% di copertura il centro sta perdendo clienti senza accorgersene."),
      ("", ""),
      ("ASSISTENZA", "351 846 6025"),
    ]
    r = 5
    for etichetta, testo in righe:
        if etichetta:
            ws.cell(r, 2, etichetta).font = A(bold=True, color=BLU)
        if testo:
            c = ws.cell(r, 3, testo); c.font = A()
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=10)
        r += 1
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 17
    for col in "CDEFGHIJ":
        ws.column_dimensions[col].width = 12
    return ws


def impostazioni(wb):
    ws = wb.create_sheet("IMPOSTAZIONI")
    ws.sheet_view.showGridLines = False
    titolo(ws, "IMPOSTAZIONI", "I sei numeri che comandano tutti gli avvisi. Giallo = si scrive.", "D")
    voci = [("Sedute per ciclo", 12, "Arrivata a questo numero scatta ULTIMA SEDUTA."),
            ("Rientro da (giorni)", 21, "Prima, la ricrescita non è pronta."),
            ("Rientro entro (giorni)", 28, "Oltre, la cliente è FUORI FINESTRA."),
            ("Conferma quanti giorni prima", 7, "La prima telefonata di conferma. La seconda è sempre il giorno prima."),
            ("Mantenimento ogni (giorni)", 90, "Ogni quanto si richiama chi ha finito il ciclo."),
            ("Dormiente dopo (mesi)", 6, "In MESI. Dopo questo silenzio si passa alle offerte di recupero.")]
    for i, (etichetta, valore, nota) in enumerate(voci):
        r = 5 + i
        ws.cell(r, 2, etichetta).font = A(bold=True)
        c = ws.cell(r, 3, valore); c.font = A(bold=True, color=BLU); c.fill = GIALLO
        c.alignment = Alignment(horizontal="center")
        ws.cell(r, 4, nota).font = A(color="FF5A6B79")
    ws["E3"] = "L'ordine di lavoro"; ws["E3"].font = A(bold=True, color=BLU)
    ws["E4"] = "n."; ws["F4"] = "avviso"
    for c in ("E4", "F4"):
        ws[c].font = Font(name="Arial", size=9, bold=True, color="FFFFFFFF"); ws[c].fill = TESTATA
    for i, a in enumerate(AVVISI):
        r = 5 + i
        ws.cell(r, 5, i + 1).font = A(bold=True, color=BLU)
        ws.cell(r, 5).alignment = Alignment(horizontal="center")
        cel = ws.cell(r, 6, a); cel.font = A(); cel.fill = GIALLO
    ws["F14"] = "La posizione in questa lista È la priorità: sposti una riga e tutto il centro si riordina."
    ws["F14"].font = A(italic=True, color="FF5A6B79")
    ws["F15"] = "Prima quello che oggi si spegne da solo, poi quello che resta acceso finché non lo chiudi."
    ws["F15"].font = A(italic=True, color="FF5A6B79")
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 30; ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 58; ws.column_dimensions["E"].width = 5
    ws.column_dimensions["F"].width = 26
    return ws


def incentivi(wb):
    ws = wb.create_sheet("INCENTIVI")
    ws.sheet_view.showGridLines = False
    titolo(ws, "INCENTIVI", "Tre offerte, una per famiglia. Il nome finisce da solo nella rubrica.", "E")
    for col, testo, largo in [("B", "Famiglia", 16), ("C", "Nome dell'offerta", 22),
                              ("D", "Cosa dice — finisce nel messaggio", 52), ("E", "Valida fino al", 14)]:
        c = ws[col + "4"]; c.value = testo
        c.font = Font(name="Arial", size=9, bold=True, color="FFFFFFFF"); c.fill = TESTATA
        ws.column_dimensions[col].width = largo
    righe = [("In ciclo", "Offerta Laser 1", "Per i rientri e per chi apre una zona nuova."),
             ("Mantenimento", "Offerta Laser 2", "Il richiamo a 90 giorni: deve costare poco e fissarsi subito."),
             ("Persa", "Offerta Laser 3", "Recupero delle addormentate: l'unica che può essere aggressiva.")]
    for i, (fam, nome, testo) in enumerate(righe):
        r = 5 + i
        ws.cell(r, 2, fam).font = A(bold=True)
        for col, v in ((3, nome), (4, testo), (5, None)):
            c = ws.cell(r, col, v); c.font = A(); c.fill = GIALLO; c.border = SOTTILE
        ws.cell(r, 5).number_format = "dd/mm/yyyy"
    ws["B9"] = "La famiglia non si tocca: è quella che collega l'offerta all'avviso."
    ws["B9"].font = A(italic=True, color="FF5A6B79")
    ws.column_dimensions["A"].width = 2
    return ws


def clienti(wb, righe_esempio):
    ws = wb.create_sheet("CLIENTI")
    ws.sheet_view.showGridLines = False
    titolo(ws, "CLIENTI — la rubrica", "Giallo = lo scrivi tu · Grigio = si calcola da solo", "F")

    for col, testa, largo, chi, fmt in COLONNE:
        ws.column_dimensions[col].width = largo
        c = ws[col + "4"]; c.value = testa
        c.font = Font(name="Arial", size=9, bold=True, color="FFFFFFFF")
        c.fill = TESTATA
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 30

    calcolate = {"B": "B5", "F": "F5", "J": "J5", "K": "K5", "L": "L5",
                 "M": "M5", "O": "O5", "P": "P5", "S": "S5", "T": "T5"}

    for r in range(PRIMA, ULTIMA + 1):
        for col, testa, largo, chi, fmt in COLONNE:
            c = ws[col + str(r)]
            c.font = A(); c.border = SOTTILE
            if fmt: c.number_format = fmt
            if chi == "formula":
                c.value = perRiga(calcolate[col], r)
                c.fill = CALC
            else:
                c.fill = GIALLO
        ws[("B%d" % r)].font = A(bold=True)

    # una riga di esempio, per far vedere il formato
    for i, dati in enumerate(righe_esempio):
        r = PRIMA + i
        for col, v in dati.items():
            ws[col + str(r)] = v

    # le tendine
    v_stato = DataValidation(type="list", formula1='"%s"' % ",".join(STATI), allow_blank=True,
                             errorTitle="Stato non valido", error="Scegli una voce dalla tendina.")
    v_data = DataValidation(type="date", operator="between",
                            formula1="DATE(2015,1,1)", formula2="DATE(2100,1,1)", allow_blank=True,
                            errorTitle="Serve una data", error="Scrivi una data, non del testo.")
    v_sed = DataValidation(type="whole", operator="between", formula1="0", formula2="99", allow_blank=True,
                           errorTitle="Serve un numero", error="Le sedute fatte sono un numero da 0 a 99.")
    v_prezzo = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                              errorTitle="Serve un importo", error="Il prezzo a seduta è un numero.")
    for v, rif in ((v_stato, "E"), (v_data, "C"), (v_data, "I"), (v_data, "Q"), (v_sed, "D"), (v_prezzo, "N")):
        if v not in ws.data_validations.dataValidation:
            ws.add_data_validation(v)
        v.add("%s%d:%s%d" % (rif, PRIMA, rif, ULTIMA))

    # i colori dell'avviso, sulla riga intera
    campo = "A%d:R%d" % (PRIMA, ULTIMA)
    for chiave, riempi, testo in (("FC_ROSSO", ROSSO_F, ROSSO_T),
                                  ("FC_GIALLO", GIALL_F, GIALL_T),
                                  ("FC_VERDE", VERDE_F, VERDE_T)):
        ws.conditional_formatting.add(campo, FormulaRule(formula=[perExcel(chiave)[1:]], fill=riempi, stopIfTrue=False))

    ws.column_dimensions["S"].hidden = True
    ws.column_dimensions["T"].hidden = True
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = "A4:R%d" % ULTIMA

    ws["B4"].comment = Comment(
        "Non scrivere qui dentro: è una formula.\n"
        "Il numero davanti è la priorità, così ordinando dalla A alla Z\n"
        "le telefonate vengono già nell'ordine giusto.", "NIPEC", height=110, width=330)
    ws["S4"].comment = Comment(
        "Colonna di servizio: qui sta il ragionamento che decide l'avviso.\n"
        "È nascosta apposta. Non cancellarla.", "NIPEC", height=90, width=320)
    return ws


def riepilogo(wb):
    ws = wb.create_sheet("RIEPILOGO")
    ws.sheet_view.showGridLines = False
    titolo(ws, "RIEPILOGO", "Le celle C5, C13, C17, C20, C26 e F5 le legge il master NIPEC: non spostarle.", "G")

    def dato(r, etichetta, chiave, fmt=None, nota=None):
        ws.cell(r, 2, etichetta).font = A(bold=True)
        c = ws.cell(r, 3, perExcel(chiave)); c.font = Font(name="Arial", size=12, bold=True, color=BLU)
        c.fill = CALC; c.alignment = Alignment(horizontal="center")
        if fmt: c.number_format = fmt
        if nota: ws.cell(r, 4, nota).font = A(color="FF5A6B79")

    dato(5,  "DA RICHIAMARE",              "R_C5",  "0",          "quante righe hanno un avviso acceso adesso")
    dato(7,  "Copertura appuntamento",     "R_C7",  "0%",         "sotto il 90% il centro sta perdendo clienti")
    dato(8,  "Dormienti",                  "R_C8",  "0",          "ferme da oltre i mesi impostati")
    dato(13, "Percorsi attivi",            "R_C13", "0",          "clienti in ciclo")
    dato(17, "Persone in rubrica",         "R_C17", "0",          "schede totali")
    dato(20, "Nuove negli ultimi 30 gg",   "R_C20", "0",          "contatti entrati questo mese")
    dato(26, "Ultimo inserimento",         "R_C26", "dd/mm/yyyy", "l'ultima volta che qualcuno ha toccato la rubrica")

    ws["E5"] = "Da incassare €";        ws["E6"] = "Già lasciati €";   ws["E7"] = "Media per cliente €"
    for r, chiave in ((5, "R_F5"), (6, "R_F6"), (7, "R_F7")):
        ws.cell(r, 5).font = A(bold=True)
        c = ws.cell(r, 6, perExcel(chiave)); c.font = Font(name="Arial", size=12, bold=True, color=BLU)
        c.fill = CALC; c.number_format = '#,##0 "€"'; c.alignment = Alignment(horizontal="center")

    ws["B29"] = "Quante per avviso"; ws["B29"].font = A(bold=True, color=BLU)
    for i, a in enumerate(AVVISI):
        r = 30 + i
        c = ws.cell(r, 2, "=IMPOSTAZIONI!F%d" % (5 + i)); c.font = A(); c.fill = CALC
        q = ws.cell(r, 3, perRiga("R_C30", r)); q.font = A(bold=True); q.fill = CALC
        q.alignment = Alignment(horizontal="center"); q.number_format = "0"

    ws["B39"] = "In Excel alcune celle restano vuote finché non ricalcola: è normale, non è rotto."
    ws["B39"].font = A(italic=True, color="FF5A6B79")
    ws.column_dimensions["A"].width = 2;  ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 12; ws.column_dimensions["D"].width = 48
    ws.column_dimensions["E"].width = 20; ws.column_dimensions["F"].width = 14
    return ws


def oggi(wb):
    ws = wb.create_sheet("OGGI")
    ws.sheet_view.showGridLines = False
    titolo(ws, "OGGI", "Solo chi va chiamato, già in ordine. Si guarda questo, la mattina.", "F")
    teste = [("A", "Nome", 24), ("B", "COSA FARE OGGI", 26), ("C", "Prossimo", 13),
             ("D", "Sedute", 9), ("E", "Stato", 15), ("F", "Promemoria", 40), ("G", "chiave", 10)]
    for col, testa, largo in teste:
        ws.column_dimensions[col].width = largo
        c = ws[col + "4"]; c.value = testa
        c.font = Font(name="Arial", size=9, bold=True, color="FFFFFFFF"); c.fill = TESTATA
    for i in range(30):
        r = 5 + i
        ws.cell(r, 7, perRiga("O_G5", r)).font = A()
        for n, col in enumerate("ABCDEF"):
            f = perRiga("O_A5", r).replace("CLIENTI!A$5:A$304", "CLIENTI!%s$5:%s$304" % (col, col))
            c = ws.cell(r, n + 1, f); c.font = A(); c.border = SOTTILE
            if col == "B": c.font = A(bold=True)
            if col == "C": c.number_format = "dd/mm/yyyy"
    for chiave, riempi, testo in (("FC_ROSSO", ROSSO_F, ROSSO_T), ("FC_GIALLO", GIALL_F, GIALL_T),
                                  ("FC_VERDE", VERDE_F, VERDE_T)):
        ws.conditional_formatting.add("A5:F34", FormulaRule(formula=[perExcel(chiave)[1:]], fill=riempi))
    ws.column_dimensions["G"].hidden = True
    ws.freeze_panes = "A5"
    ws["A36"] = "Se qui non c'è nessuno, oggi non deve chiamare nessuno."
    ws["A36"].font = A(italic=True, color="FF5A6B79")
    return ws


def costruisci(percorso, esempi):
    wb = Workbook()
    wb.remove(wb.active)
    leggimi(wb); oggi(wb); clienti(wb, esempi); riepilogo(wb); incentivi(wb); impostazioni(wb)
    # Excel non trova valori in cache: cosi' ricalcola da solo appena si apre
    wb.calculation.fullCalcOnLoad = True
    wb.properties.title = "Cruscotto Clienti · laser"
    wb.properties.creator = "NIPEC · Tecnologie Estetiche"
    wb.save(percorso)
    return percorso
